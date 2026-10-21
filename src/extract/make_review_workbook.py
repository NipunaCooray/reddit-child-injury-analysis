#!/usr/bin/env python3
"""
Build the hand-review workbook for a Phase 2 pilot (see src/extract/HAND_REVIEW.md).

One primary run is reviewed cell by cell; the other runs are shown only where they
disagree with it, because that is where definitions are ambiguous. Showing them
everywhere would anchor the reviewer on model output.

Sheets: Instructions, Posts, Review, Codebook, Decision log, Summary (live formulas).
"""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from common import find_repo_root, load_jsonl
from schema import (FIELDS, VOCAB, DEFINITIONS, VALUE_NOTES, SCHEMA_VERSION, AUDIT_FIELD,
                    CONDITIONAL_FALL_FIELDS, field_type)

VERDICTS = ["correct", "partially_correct", "wrong_value", "over_inferred", "missed", "wrong_status", "span_mismatch", "unsure"]
VERDICT_HELP = {
    "correct": "Right value and status; the span supports it (includes correctly not_stated cells).",
    "partially_correct": "Multi-valued or free-text cell that is part right (e.g. one of two help sources). "
                         "Put the full right answer in 'Correct value'. Counts as not correct in accuracy.",
    "wrong_value": "The post supports a different value. Put the right one in 'Correct value'.",
    "over_inferred": "Model filled it but the post does not support it; should be not_stated.",
    "missed": "Model said not_stated but the post clearly answers it. Put the value in 'Correct value'.",
    "wrong_status": "Right value, but stated vs inferred is wrong.",
    "span_mismatch": "The quote is real but does not support the value.",
    "unsure": "Cannot decide; explain in Note and add to the Decision log.",
}
# Fields checked in full on every post (tier 1), whatever the models did
PRIORITY_FIELDS = {"supervision_status", AUDIT_FIELD, "mechanism_iceci", "care_pathway_outcome",
                   "witnessed", *CONDITIONAL_FALL_FIELDS}
# Questions the pilot raised but did not settle; seeded into the Decision log as OPEN
OPEN_QUESTIONS = [
    (AUDIT_FIELD, "Child causes an injury to an adult; child unhurt (e.g. 'My toddler broke my nose')."),
    (AUDIT_FIELD, "Injury during pregnancy; baby reported fine (e.g. car crash at 9 weeks)."),
    (AUDIT_FIELD, "Near-miss with no injury (e.g. slipped out of a pool float, pebble in mouth)."),
    (AUDIT_FIELD, "Allergic reaction needing ED care."),
    (AUDIT_FIELD, "Insect bites or stings with no further harm (e.g. mosquito bites)."),
    (AUDIT_FIELD, "Injury mentioned only as backstory, not the subject of the post."),
    ("fall_height", "Fall height rarely given even in fall posts: keep, mark rare_by_design, or cut?"),
]

# Study guide and codebook (purpose, design, coding notes, worked example); readers need access to it
STUDY_GUIDE_URL = "https://claude.ai/artifact/AmnHgTaUPuqYjf5fYBct3W"

FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_MUTED = Font(name=FONT, size=9, italic=True, color="595959")
FILL_HEAD = PatternFill("solid", fgColor="404040")
FILL_INPUT = PatternFill("solid", fgColor="FFFF99")
FILL_TIER1 = PatternFill("solid", fgColor="FCE4D6")
FILL_POST = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

def val_str(v):
    if v is None:
        return ""
    return "|".join(v) if isinstance(v, list) else str(v)

def header(ws, row, cols, widths=None):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, WRAP
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

def est_height(text, width_chars, line_pt=13, cap=409):
    lines = sum(max(1, -(-len(p) // width_chars)) for p in (text or "").split("\n"))
    return min(cap, max(15, lines * line_pt))

def chunk_body(text, limit=3800):
    # Excel rows top out at 409pt, so long posts continue on the next row(s)
    paras, chunks, cur = (text or "").split("\n"), [], ""
    for p in paras:
        if cur and len(cur) + len(p) + 1 > limit:
            chunks.append(cur)
            cur = ""
        while len(p) > limit:
            chunks.append(p[:limit])
            p = p[limit:]
        cur = f"{cur}\n{p}" if cur else p
    if cur or not chunks:
        chunks.append(cur)
    return chunks

def main():
    repo_root = find_repo_root(Path(__file__).resolve())

    ap = argparse.ArgumentParser(description="Build the hand-review workbook for a Phase 2 pilot.")
    ap.add_argument("--run", action="append", required=True, metavar="TAG=PATH",
                    help="An *_extractions.jsonl (repeatable). The --primary run is reviewed; others shown on disagreement.")
    ap.add_argument("--primary", required=True, help="TAG of the run to review")
    ap.add_argument("--out", default=None, help="Default: reports/extraction_qa/hand_review/hand_review_<primary>.xlsx")
    args = ap.parse_args()

    runs = {}
    for item in args.run:
        tag, _, path = item.partition("=")
        runs[tag] = {r["id"]: r for r in load_jsonl(Path(path).expanduser().resolve())}
    if args.primary not in runs:
        raise SystemExit(f"--primary {args.primary!r} is not one of the --run tags {list(runs)}")
    others = [t for t in runs if t != args.primary]
    prim = runs[args.primary]
    ids = [i for i in prim if all(i in runs[t] for t in others)]
    if len(ids) != len(prim):
        print(f"[warn] {len(prim) - len(ids)} primary posts missing from another run; skipped")
    versions = {r.get("extraction_schema_version") for r in prim.values()}
    model = next(iter(prim.values())).get("model")

    out_path = Path(args.out).expanduser().resolve() if args.out else \
        repo_root / "reports" / "extraction_qa" / "hand_review" / f"hand_review_{args.primary}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # ---------------- Instructions ----------------
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    rows = [
        ("Phase 2 hand review", None),
        ("Model reviewed", f"{model} (run tag '{args.primary}'); others shown on disagreement: {', '.join(others) or 'none'}"),
        ("Schema version", ", ".join(sorted(v for v in versions if v)) + f"  (workbook built against schema {SCHEMA_VERSION})"),
        ("Posts / cells", f"{len(ids)} posts x {len(FIELDS)} fields = {len(ids) * len(FIELDS)} review rows"),
        ("Reviewer", ""),
        ("Read first", f"Study guide and codebook (what the study is, the status decision, coding notes, a worked example): "
                       f"{STUDY_GUIDE_URL}  |  Protocol for this review: src/extract/HAND_REVIEW.md in the repo."),
        ("Data handling", "This file contains full post text about injured children. Keep it on your own machine; do not "
                          "commit, email or upload it unless the study lead says otherwise."),
        ("Purpose", "Check the model against the post to (1) lock the schema, (2) confirm the model before the full run, "
                    "(3) settle the audit edge cases. This is NOT the formal reliability study: that comes later on a fresh "
                    "random sample, coded blind by two people. Full protocol: src/extract/HAND_REVIEW.md"),
        ("Yellow cells", "Only edit yellow cells: Reviewer (above), Review!N:P, Decision log!D:H, Posts!F. Everything else is generated."),
        ("Step 1", "Open 'Posts', read post N in full BEFORE looking at any model output. Note what happened, who was watching, "
                   "where, what help was sought, and whether this is really an injury."),
        ("Step 2", "Open 'Review', filter column A to post N. For each row, judge the model's value + status + span against the post "
                   "and pick a Verdict. 'Codebook' has each field's definition and allowed values; the study guide's coding "
                   "notes settle the common hard cases."),
        ("Step 3", "Columns L-M show other models' values ONLY where they disagree. Use them to ask: is the model wrong, or is the "
                   "definition ambiguous? If ambiguous, write a rule in 'Decision log'."),
        ("Step 4", "Wrong or missed: put the right value in 'Correct value' (enum fields offer a dropdown; for help_seeking type "
                   "values joined with |). Explain anything non-obvious in 'Note'."),
        ("Tier", "Tier 1 rows (shaded) must be reviewed on every post: priority fields (supervision, injury check, mechanism, care "
                 "pathway, witnessed, fall fields) and any field where models disagree. Tier 2 = quick check. If time is short, "
                 "finish Tier 1 on all posts before Tier 2 on any."),
        ("Pacing", "Do posts 1-10 first and write Decision log rules as you go; most rules come from the first batch. Then the rest in "
                   "2-3 sittings. Expect ~5-6 min per post."),
        ("Not_stated cells", "A not_stated cell is 'correct' if the post really does not say. It is 'missed' if the post plainly answers it. "
                             "Prefer not_stated over a guess: do not credit inferences from typical parenting context."),
        ("Phase 1 note", "Posts!E is the Phase 1 MODEL's own note, not a human judgement. Do not treat it as ground truth."),
    ]
    for r, (k, v) in enumerate(rows, 1):
        a = ws.cell(row=r, column=1, value=k)
        b = ws.cell(row=r, column=2, value=v)
        a.font, b.font = (F_TITLE if r == 1 else F_BOLD), F_BODY
        a.alignment = b.alignment = WRAP
        if k == "Reviewer":
            b.fill = FILL_INPUT
    r = len(rows) + 2
    ws.cell(row=r, column=1, value="Verdict codes").font = F_BOLD
    for v in VERDICTS:
        r += 1
        ws.cell(row=r, column=1, value=v).font = F_BOLD
        ws.cell(row=r, column=2, value=VERDICT_HELP[v]).font = F_BODY
    r += 2
    ws.cell(row=r, column=1, value="EXAMPLE (format only)").font = F_BOLD
    ws.cell(row=r, column=2, value="A filled-in Review row for an invented post, 'Toddler fell off the changing table'. "
                                    "Not part of the data.").font = F_MUTED
    example = [
        ("Field", "supervision_status"),
        ("Model value / status", "direct_supervision / inferred"),
        ("Model span", "I turned around to grab a wipe and she rolled right off"),
        ("Verdict", "wrong_value"),
        ("Correct value", "proximate_supervision"),
        ("Note", "Caregiver within reach but looking away at that moment = proximate, per Codebook gloss."),
    ]
    for k, v in example:
        r += 1
        ws.cell(row=r, column=1, value=k).font = F_MUTED
        c = ws.cell(row=r, column=2, value=v)
        c.font = F_MUTED
        if k in ("Verdict", "Correct value", "Note"):
            c.fill = FILL_INPUT

    # ---------------- Posts ----------------
    wp = wb.create_sheet("Posts")
    header(wp, 1, ["Post #", "Post ID", "Title", "Body", "Phase 1 model note (NOT ground truth)", "Read? (x)", "Permalink"],
           [7, 10, 30, 110, 30, 8, 30])
    wp.freeze_panes = "A2"
    row = 2
    post_no = {}
    for n, i in enumerate(ids, 1):
        post_no[i] = n
        rec = prim[i]
        p1 = rec.get("phase1_labels") or {}
        chunks = chunk_body(rec.get("selftext") or "")
        for k, chunk in enumerate(chunks):
            cont = k > 0
            vals = [n, i, (rec["title"] if not cont else f"(continued {k + 1}/{len(chunks)})"), chunk,
                    (p1.get("rationale_short", "") if not cont else ""), None,
                    (rec.get("permalink", "") if not cont else "")]
            for col, v in enumerate(vals, 1):
                c = wp.cell(row=row, column=col, value=v)
                c.font, c.alignment = F_BODY, WRAP
                if n % 2 == 0:
                    c.fill = FILL_POST
            if not cont:
                wp.cell(row=row, column=3).font = F_BOLD
                wp.cell(row=row, column=6).fill = FILL_INPUT
            wp.row_dimensions[row].height = est_height(chunk, 105)
            row += 1

    # ---------------- Codebook (+ named value lists for the Correct value dropdown) ----------------
    wc = wb.create_sheet("Codebook")
    header(wc, 1, ["Phase", "Cell", "Field", "Type", "Definition", "Allowed values", "Value notes"],
           [12, 18, 24, 10, 70, 45, 55])
    wc.freeze_panes = "A2"
    for r, (phase, cell, f) in enumerate(FIELDS, 2):
        notes = "\n".join(f"{k} = {v}" for k, v in VALUE_NOTES.get(f, {}).items())
        vals = [phase, cell, f, field_type(f), DEFINITIONS[f], "\n".join(VOCAB.get(f, ())) or "free text", notes]
        for col, v in enumerate(vals, 1):
            c = wc.cell(row=r, column=col, value=v)
            c.font, c.alignment = F_BODY, WRAP
        wc.cell(row=r, column=3).font = F_BOLD
    wv = wb.create_sheet("Vocab")
    for col, (f, values) in enumerate(VOCAB.items(), 1):
        wv.cell(row=1, column=col, value=f).font = F_BOLD
        for r, v in enumerate(values, 2):
            wv.cell(row=r, column=col, value=v).font = F_BODY
        letter = get_column_letter(col)
        wb.defined_names[f] = DefinedName(f, attr_text=f"Vocab!${letter}$2:${letter}${len(values) + 1}")
        wv.column_dimensions[letter].width = 28
    wv.sheet_state = "hidden"

    # ---------------- Review ----------------
    wr = wb.create_sheet("Review", 2)
    # layout is fixed at two comparison columns (L, M) so the Summary formulas never move
    compare_cols = [f"{t} value (if differs)" for t in others[:2]]
    compare_cols += ["(no other run)"] * (2 - len(compare_cols))
    cols = ["Post #", "Post ID", "Title", "Field", "Phase", "Tier",
            f"{args.primary} value", f"{args.primary} status", f"{args.primary} span", "Guardrail flags",
            "Models agree?"] + compare_cols + ["Verdict", "Correct value", "Note"]
    header(wr, 1, cols, [7, 10, 26, 22, 11, 6, 24, 10, 45, 18, 9, 24, 24, 15, 24, 40])
    wr.freeze_panes = "G2"
    first, row = 2, 2
    for i in ids:
        rp = prim[i]
        for phase, _, f in FIELDS:
            pc = rp["fields"][f]
            pv = val_str(pc["value"])
            flags = sorted({v["code"] for v in rp["violations"] if v["field"] == f})
            diffs, agree = [], True
            for t in others[:2]:
                oc = runs[t][i]["fields"][f]
                ov = val_str(oc["value"])
                if field_type(f) == "text":
                    differs = (oc["status"] == "not_stated") != (pc["status"] == "not_stated")
                else:
                    differs = ov != pv
                agree &= not differs
                diffs.append((ov or "(not_stated)") if differs else "")
            tier = 1 if (f in PRIORITY_FIELDS or not agree) else 2
            vals = [post_no[i], i, rp["title"][:60], f, phase, tier, pv or "(not_stated)", pc["status"],
                    pc["evidence_span"] or "", "|".join(flags), "Y" if agree else "N"]
            vals += diffs + [""] * (2 - len(diffs)) + [None, None, None]
            for col, v in enumerate(vals, 1):
                c = wr.cell(row=row, column=col, value=v)
                c.font, c.alignment, c.border = F_BODY, WRAP, THIN
                if tier == 1 and col <= 13:
                    c.fill = FILL_TIER1
            for col in (14, 15, 16):
                wr.cell(row=row, column=col).fill = FILL_INPUT
            wr.cell(row=row, column=4).font = F_BOLD
            row += 1
    last = row - 1
    wr.auto_filter.ref = f"A1:P{last}"

    dv = DataValidation(type="list", formula1=f'"{",".join(VERDICTS)}"', allow_blank=True)
    dv.error, dv.errorTitle = "Pick a verdict from the list.", "Verdict"
    wr.add_data_validation(dv)
    dv.add(f"N{first}:N{last}")
    # Correct value: dropdown of the field's allowed values via its named range; free entry
    # still allowed (alerts off) for free-text fields and multi-valued help_seeking
    dv2 = DataValidation(type="list", formula1=f"INDIRECT($D{first})", allow_blank=True, showErrorMessage=False)
    wr.add_data_validation(dv2)
    dv2.add(f"O{first}:O{last}")
    wr["N1"].comment = Comment("Pick one: " + ", ".join(VERDICTS) + ". See Instructions for meanings.", "generator")

    # ---------------- Decision log ----------------
    wd = wb.create_sheet("Decision log")
    header(wd, 1, ["#", "Status", "Field", "Situation (what was unclear)", "Rule decided", "Example post ID",
                   "Schema change needed? (Y/N)", "Notes"], [5, 10, 22, 60, 50, 14, 14, 40])
    wd.freeze_panes = "A2"
    seed = [("EXAMPLE", "supervision_status", "Caregiver in the same room but cooking, not watching, when the child fell.",
             "proximate_supervision (nearby but not watching at that moment).", "", "N",
             "Format example only. Applies the existing Codebook gloss.")]
    seed += [("OPEN", f, q, None, None, None, None) for f, q in OPEN_QUESTIONS]
    for r, (status, f, situation, rule, pid, change, note) in enumerate(seed, 2):
        vals = [r - 1, status, f, situation, rule, pid, change, note]
        for col, v in enumerate(vals, 1):
            c = wd.cell(row=r, column=col, value=v)
            c.font = F_MUTED if status == "EXAMPLE" else F_BODY
            c.alignment = WRAP
            if col >= 4 and status != "EXAMPLE":
                c.fill = FILL_INPUT
    for r in range(len(seed) + 2, len(seed) + 32):
        wd.cell(row=r, column=1, value=r - 1).font = F_BODY
        for col in range(2, 9):
            wd.cell(row=r, column=col).fill = FILL_INPUT

    # ---------------- Summary (live formulas over Review) ----------------
    wsu = wb.create_sheet("Summary", 1)
    rng = lambda col: f"Review!${col}${first}:${col}${last}"
    wsu["A1"], wsu["A1"].font = "Review progress and model accuracy (updates as you fill Review!N)", F_TITLE
    wsu["A3"], wsu["B3"] = "Rows reviewed", f'=COUNTIFS({rng("N")},"<>")'
    wsu["A4"], wsu["B4"] = "Total rows", f'=COUNTIFS({rng("D")},"<>")'
    wsu["A5"], wsu["B5"] = "Tier 1 reviewed / total", \
        f'=COUNTIFS({rng("F")},1,{rng("N")},"<>")&" / "&COUNTIF({rng("F")},1)'
    wsu["A6"], wsu["B6"] = "Progress", "=IF(B4=0,0,B3/B4)"
    wsu["B6"].number_format = "0.0%"
    for r in range(3, 7):
        wsu.cell(row=r, column=1).font = F_BOLD
        wsu.cell(row=r, column=2).font = F_BODY
    hdr = ["Field", "Tier 1 field?", "Reviewed"] + VERDICTS + ["Accuracy (correct / decided)"]
    header(wsu, 8, hdr, [24, 10, 10] + [11] * len(VERDICTS) + [16])
    for r, (_, _, f) in enumerate(FIELDS, 9):
        wsu.cell(row=r, column=1, value=f).font = F_BOLD
        wsu.cell(row=r, column=2, value="Y" if f in PRIORITY_FIELDS else "").font = F_BODY
        wsu.cell(row=r, column=3, value=f'=COUNTIFS({rng("D")},$A{r},{rng("N")},"<>")').font = F_BODY
        for k, v in enumerate(VERDICTS):
            col = 4 + k
            wsu.cell(row=r, column=col, value=f'=COUNTIFS({rng("D")},$A{r},{rng("N")},"{v}")').font = F_BODY
        unsure_col = get_column_letter(4 + VERDICTS.index("unsure"))
        acc = wsu.cell(row=r, column=4 + len(VERDICTS),
                       value=f'=IF(C{r}-{unsure_col}{r}=0,"",D{r}/(C{r}-{unsure_col}{r}))')
        acc.font, acc.number_format = F_BODY, "0.0%"
    tr = 9 + len(FIELDS)
    wsu.cell(row=tr, column=1, value="ALL FIELDS").font = F_BOLD
    for col in range(3, 4 + len(VERDICTS)):
        L = get_column_letter(col)
        wsu.cell(row=tr, column=col, value=f"=SUM({L}9:{L}{tr - 1})").font = F_BOLD
    unsure_col = get_column_letter(4 + VERDICTS.index("unsure"))
    acc = wsu.cell(row=tr, column=4 + len(VERDICTS), value=f'=IF(C{tr}-{unsure_col}{tr}=0,"",D{tr}/(C{tr}-{unsure_col}{tr}))')
    acc.font, acc.number_format = F_BOLD, "0.0%"
    wsu.cell(row=tr + 2, column=1, value="Accuracy excludes 'unsure'. over_inferred + missed are the rates that decide model "
                                          "choice and which fields need redefinition.").font = F_MUTED
    wsu.freeze_panes = "B9"

    # openpyxl writes formulas without cached values; make Excel compute them on open
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f"Wrote {out_path}\n  {len(ids)} posts, {last - first + 1} review rows "
          f"({sum(1 for r in range(first, last + 1) if wr.cell(row=r, column=6).value == 1)} tier 1)")

if __name__ == "__main__":
    main()
